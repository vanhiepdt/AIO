import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document


class WordEditor:
    """Simple GUI application to replace placeholders in a Word document."""

    def __init__(self, master: tk.Tk) -> None:
        self.master = master
        master.title("Word Placeholder Replacer")
        self.doc = None

        self.open_button = tk.Button(master, text="Open Document", command=self.open_file)
        self.open_button.pack(pady=5)

        self.placeholder_label = tk.Label(master, text="Placeholder:")
        self.placeholder_label.pack()
        self.placeholder_entry = tk.Entry(master, width=40)
        self.placeholder_entry.pack()

        self.replacement_label = tk.Label(master, text="Replacement:")
        self.replacement_label.pack()
        self.replacement_entry = tk.Entry(master, width=40)
        self.replacement_entry.pack()

        self.replace_button = tk.Button(master, text="Replace", command=self.replace_text)
        self.replace_button.pack(pady=5)

        self.save_button = tk.Button(master, text="Save As", command=self.save_file)
        self.save_button.pack()

    def open_file(self) -> None:
        """Load a .docx file."""
        file_path = filedialog.askopenfilename(filetypes=[("Word documents", "*.docx")])
        if file_path:
            self.doc = Document(file_path)
            messagebox.showinfo("Loaded", f"Loaded document: {file_path}")

    def replace_text(self) -> None:
        """Replace placeholder text in the loaded document."""
        if not self.doc:
            messagebox.showwarning("No document", "Please open a document first.")
            return
        placeholder = self.placeholder_entry.get()
        replacement = self.replacement_entry.get()
        if not placeholder:
            messagebox.showwarning("No placeholder", "Please enter a placeholder.")
            return

        for paragraph in self.doc.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace(placeholder, replacement)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.text = run.text.replace(placeholder, replacement)

        messagebox.showinfo("Replaced", "Replacement complete.")

    def save_file(self) -> None:
        """Save the modified document."""
        if not self.doc:
            messagebox.showwarning("No document", "Please open a document first.")
            return
        save_path = filedialog.asksaveasfilename(defaultextension=".docx",
                                                filetypes=[("Word documents", "*.docx")])
        if save_path:
            self.doc.save(save_path)
            messagebox.showinfo("Saved", f"Document saved as: {save_path}")


def main() -> None:
    root = tk.Tk()
    app = WordEditor(root)
    root.mainloop()


if __name__ == "__main__":
    main()
